"""A downloaded or uploaded ``.xlsx`` → ``Grid`` (dev, tests, the Admin upload; SPEC §4a).

``openpyxl`` is imported inside ``read_xlsx``, so it stays off the cold-start
path: only the Admin upload and dev runs pay for it. A workbook is opened
twice, once for formulas and once for the values the sheet app last computed.
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
import zlib
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError

from fha.sources.league_sheet.grid import Cell, Grid, Tab, Value
from fha.sources.league_sheet.models import LeagueSheetError


def read_xlsx(source: bytes | str | Path) -> Grid:
    import openpyxl  # lazily: see the module docstring
    from openpyxl.utils.exceptions import InvalidFileException

    # ParseError: a valid zip holding malformed XML.
    unreadable = (
        OSError,
        KeyError,
        ValueError,
        ParseError,
        zipfile.BadZipFile,
        InvalidFileException,
        zlib.error,  # corrupt deflate data: a damaged upload
        TypeError,  # malformed document properties
        NotImplementedError,  # an unsupported zip compression method
        EOFError,
    )

    def load(data_only: bool) -> Any:
        stream = io.BytesIO(source) if isinstance(source, bytes) else source
        try:
            return openpyxl.load_workbook(stream, data_only=data_only, read_only=False)
        except unreadable as e:
            raise LeagueSheetError(f"not a readable .xlsx file ({type(e).__name__})") from None

    formulas, values = load(data_only=False), load(data_only=True)
    tabs = []
    for sheet in formulas.worksheets:
        computed = values[sheet.title]
        cells: dict[tuple[int, int], Cell] = {}
        for row in sheet.iter_rows():
            for c in row:
                raw = c.value
                if raw is None:
                    continue
                formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                if formula is None and not isinstance(raw, (str, int, float, bool, *TIMES)):
                    formula = str(getattr(raw, "text", raw))  # e.g. an array formula
                value = computed.cell(row=c.row, column=c.column).value
                cells[(c.row, c.column)] = Cell(_value(value), formula)
        tabs.append(Tab(sheet.title, cells))
    return Grid(tuple(tabs))


TIMES = (dt.date, dt.time, dt.timedelta)  # plain values openpyxl gives as objects


def _value(raw: Any) -> Value:
    if raw is None or isinstance(raw, str | int | float | bool):
        return raw
    if isinstance(raw, dt.timedelta):
        return str(raw)
    if isinstance(raw, dt.date | dt.time):
        return raw.isoformat()
    return str(raw)  # anything else openpyxl might hand back, as text
