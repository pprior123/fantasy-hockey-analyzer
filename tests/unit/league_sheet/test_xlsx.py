"""The .xlsx grid reader: values and formulas together, from bytes or a path."""

import datetime as dt
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from fha.sources.league_sheet.grid import Cell, Grid, Tab
from fha.sources.league_sheet.models import LeagueSheetError
from fha.sources.league_sheet.parse import parse_sheet
from fha.sources.league_sheet.xlsx import read_xlsx
from tests.unit.league_sheet.sheets import CONTACT, IRRow, Player, TeamTab, grid, notes_tab, roster
from tests.unit.league_sheet.xlsx_build import to_xlsx

VARIANTS = [
    TeamTab("Aces", roster(3), below=[IRRow("IR", Player("Ike Injured")), IRRow("IR+")]),
    TeamTab("Bees", roster(4), start=6, header="in_range", payroll="sum_ref"),
    TeamTab("Cats", roster(2), payroll="cell_ref"),
    TeamTab(
        "Foxes",
        [*roster(2), Player("Ben Baker", salary="???"), None],
        salary_col="G",
        pos_col="E",
        team_col="F",
        extra_salary_cols=("H",),
    ),
]


def test_values_and_formulas_are_both_read() -> None:
    source = Grid(
        (Tab("T", {(3, 3): Cell(42, "=SUM(F7:F9)"), (7, 6): Cell(42), (1, 1): Cell("x")}),)
    )
    [tab] = read_xlsx(to_xlsx(source))
    assert tab.title == "T"
    assert tab.at("C3") == Cell(42, "=SUM(F7:F9)")
    assert tab.at("F7") == Cell(42)
    assert tab.at("A1") == Cell("x")
    assert tab.at("B2") == Cell()


def test_a_formula_the_sheet_app_never_computed_has_no_value() -> None:
    source = Grid((Tab("T", {(3, 3): Cell(42, "=SUM(F7:F9)")}),))
    [tab] = read_xlsx(to_xlsx(source, cached_values=False))
    assert tab.at("C3") == Cell(None, "=SUM(F7:F9)")


def test_every_layout_variant_parses_the_same_from_xlsx_as_from_the_grid() -> None:
    source = grid(*VARIANTS, notes_tab())
    from_file = parse_sheet(read_xlsx(to_xlsx(source)))
    assert from_file == parse_sheet(source)
    assert [t.status for t in from_file.tabs] == ["ok"] * 4
    assert all(t.payroll == t.counted_total for t in from_file.tabs)


def test_a_path_works_like_bytes(tmp_path: Path) -> None:
    path = tmp_path / "sheet.xlsx"
    path.write_bytes(to_xlsx(grid(*VARIANTS)))
    assert parse_sheet(read_xlsx(path)) == parse_sheet(read_xlsx(str(path)))
    assert len(parse_sheet(read_xlsx(path)).tabs) == 4


def test_dates_times_and_bools_become_plain_values() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "T"
    ws["A1"] = dt.datetime(2026, 9, 26, 12, 30)
    ws["A2"] = dt.date(2026, 9, 26)
    ws["A3"] = dt.time(7, 5)
    ws["A4"] = True
    path_bytes = _save(wb)
    [tab] = read_xlsx(path_bytes)
    assert tab.at("A1").value == "2026-09-26T12:30:00"
    assert tab.at("A2").value in {
        "2026-09-26",
        "2026-09-26T00:00:00",
    }  # openpyxl reads dates back as datetimes
    assert tab.at("A3").value == "07:05:00"
    assert tab.at("A4").value is True


def test_times_and_durations_are_values_not_formulas() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = dt.time(7, 5)
    ws["A2"] = dt.timedelta(hours=1, minutes=30)
    [tab] = read_xlsx(_save(wb))
    assert tab.at("A1") == Cell("07:05:00")
    assert tab.at("A2").formula is None
    assert tab.at("A2").value in {
        "1:30:00",
        "1900-01-01T01:30:00",
    }  # openpyxl may read it back as a datetime


def test_unknown_value_objects_become_text() -> None:
    from decimal import Decimal

    from fha.sources.league_sheet.xlsx import _value

    assert _value(Decimal("1.5")) == "1.5"
    assert _value(dt.timedelta(minutes=2)) == "0:02:00"


def test_an_array_formula_is_read_as_its_formula_text() -> None:
    from openpyxl.worksheet.formula import ArrayFormula

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws["C3"] = ArrayFormula("C3", "=SUM(F7:F9)")
    [tab] = read_xlsx(_save(wb))
    assert tab.at("C3").formula == "=SUM(F7:F9)"


def _save(wb: openpyxl.Workbook) -> bytes:
    import io

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.mark.parametrize("junk", [b"", b"not a zip", b"PK\x03\x04broken"])
def test_an_unreadable_file_is_a_league_sheet_error(junk: bytes) -> None:
    with pytest.raises(LeagueSheetError, match=r"not a readable \.xlsx file") as caught:
        read_xlsx(junk)
    assert caught.value.__cause__ is None


def test_malformed_xml_inside_a_valid_zip_is_a_league_sheet_error() -> None:
    import io
    import zipfile

    good = zipfile.ZipFile(io.BytesIO(_save(openpyxl.Workbook())))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in good.infolist():
            junk = item.filename in ("xl/workbook.xml", "xl/worksheets/sheet1.xml")
            z.writestr(item, b"<<not xml" if junk else good.read(item))
    with pytest.raises(
        LeagueSheetError, match=r"not a readable \.xlsx file \(ParseError\)"
    ) as caught:
        read_xlsx(out.getvalue())
    assert caught.value.__cause__ is None


def test_a_missing_path_is_a_league_sheet_error(tmp_path: Path) -> None:
    with pytest.raises(LeagueSheetError, match="FileNotFoundError"):
        read_xlsx(tmp_path / "missing.xlsx")


def test_contact_details_stay_in_the_grid() -> None:
    parsed = parse_sheet(read_xlsx(to_xlsx(grid(*VARIANTS))))
    assert [c for c in CONTACT if c in repr(parsed)] == []


def test_importing_the_sources_does_not_import_openpyxl() -> None:
    """The cold-start path (SPEC §2): openpyxl loads only when an .xlsx is read."""
    code = (
        "import sys; import fha.sources.league_sheet.source, fha.sources.league_sheet.parse; "
        "print('openpyxl' in sys.modules)"
    )
    out = subprocess.run(  # noqa: S603 - a fixed command: this interpreter, our code
        [sys.executable, "-c", code], capture_output=True, text=True, check=True
    )
    assert out.stdout.strip() == "False"
