"""Build ``.xlsx`` bytes from a synthetic ``Grid``, in test code (SPEC §8: no
committed .xlsx).

openpyxl writes a formula with an empty cached value, whereas a file saved by
Google Sheets or Excel carries the computed value too. So after saving, the
values are written into the sheet XML, which gives the reader the same
file shape it will meet in a real download.
"""

from __future__ import annotations

import io
import re
import zipfile
from xml.sax.saxutils import escape

import openpyxl

from fha.sources.league_sheet.grid import Grid, a1


def to_xlsx(grid: Grid, *, cached_values: bool = True) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    for tab in grid:
        ws = wb.create_sheet(tab.title)
        for (row, col), cell in tab.cells.items():
            ws.cell(row=row, column=col, value=cell.formula or cell.value)
    raw = io.BytesIO()
    wb.save(raw)
    if not cached_values:
        return raw.getvalue()
    return _with_cached_values(raw.getvalue(), grid)


def _with_cached_values(data: bytes, grid: Grid) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", info.filename)
            if m:
                tab = grid.tabs[int(m[1]) - 1]
                xml = content.decode()
                for (row, col), cell in tab.cells.items():
                    if cell.formula and cell.value is not None:
                        xml = _cache(xml, a1(row, col), cell.value)
                content = xml.encode()
            dst.writestr(info, content)
    return out.getvalue()


def _cache(xml: str, ref: str, value: object) -> str:
    pattern = re.compile(rf'<c r="{ref}"([^>]*)><f>(.*?)</f><v\s*/>')
    if isinstance(value, bool):
        typed, text = ' t="b"', "1" if value else "0"
    elif isinstance(value, int | float):
        typed, text = "", repr(value)
    else:
        typed, text = ' t="str"', escape(str(value))
    new, n = pattern.subn(lambda m: f'<c r="{ref}"{typed}><f>{m[2]}</f><v>{text}</v>', xml)
    assert n == 1, f"no formula cell {ref} to cache"
    return new
