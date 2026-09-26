"""scripts/extract_golden.py against a synthetic workbook built here (no real data)."""

import json
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pytest
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table

from scripts import extract_golden as eg


@dataclass
class Row:
    pid: str
    first: str
    last: str
    gp: int
    ttltst: float
    aav: int | str = 1_000_000
    stats: dict[str, int] = field(
        default_factory=lambda: {
            "G": 1,
            "A": 2,
            "PIM": 3,
            "PPG": 4,
            "PPA": 5,
            "HIT": 6,
            "BLK": 7,
            "SOG": 8,
        }
    )


DIVISORS = {
    "G": 40.0,
    "A": 60.0,
    "PIM": 100.0,
    "PPP": 30.0,
    "HIT": 250.0,
    "BLK": 160.0,
    "SOG": 260.0,
}
TOP20 = "{" + ",".join(str(i) for i in range(1, 21)) + "}"


def build(
    rows: list[Row],
    *,
    source_rows: list[int] | None = None,
    extra_source: int = 0,
    table_end: int | None = None,
) -> tuple[Workbook, Workbook]:
    """A (values, formulas) pair shaped like the owner's workbook."""
    sources = source_rows or [i + 2 for i in range(len(rows))]
    values, formulas = openpyxl.Workbook(), openpyxl.Workbook()
    for wb in (values, formulas):
        wb.active.title = eg.ANALYSIS  # type: ignore[union-attr]
        wb.create_sheet(eg.STATS)
        wb.create_sheet(eg.ALIASES)
    fa_v, fa_f = values[eg.ANALYSIS], formulas[eg.ANALYSIS]
    st_v = values[eg.STATS]
    for col, header in eg.STATS_HEADERS.items():
        st_v[f"{col}1"] = header
    for src_row in range(2, max(sources) + 1 + extra_source):
        st_v[f"A{src_row}"] = f"unused{src_row}"
    ranked = sorted((r.ttltst for r in rows), reverse=True)
    n = len(rows)
    for i, (r, src) in enumerate(zip(rows, sources, strict=True)):
        row = eg.FIRST_ROW + i
        st_v[f"A{src}"] = r.pid
        s = r.stats
        cells: dict[str, int | float | str] = {
            "A": f"{r.first} {r.last}",
            "C": "COL",
            "D": "C",
            "E": r.gp,
            "F": s["G"],
            "G": s["A"],
            "J": s["PIM"],
            "K": s["SOG"],
            "M": str(s["PPG"]),
            "N": str(s["PPA"]),
            "O": s["PPG"] + s["PPA"],
            "P": s["HIT"],
            "Q": s["BLK"],
            "S": r.aav,
            "AD": r.ttltst,
            "AE": i + 1,
            "AF": ranked[i],
            "AG": (1 - (i + 1) / n) * 100,
        }
        for col, value in cells.items():
            fa_v[f"{col}{row}"] = value
        for col, template in eg.EXPECTED_FORMULAS.items():
            fa_f[f"{col}{row}"] = eg.expected(template, row, src)
        for col, template in eg.RANKING_FORMULAS.items():
            fa_f[f"{col}{row}"] = eg.expected(template, row, row)
        fa_f[f"AE{row}"] = 1 if i == 0 else f"=AE{row - 1}+1"
    fa_v[f"A{eg.FIRST_ROW + n}"] = " "  # the workbook pads with blank names
    for cat, (col, _src, div) in eg.CATEGORIES.items():
        fa_v[f"{div}{eg.DIVISOR_ROW}"] = DIVISORS[cat]
        cell = f"{div}{eg.DIVISOR_ROW}"
        fa_f[cell] = ArrayFormula(
            cell,
            f"=AVERAGE(LARGE({col}3:{col}902,{TOP20}))/AVERAGE(LARGE(E$3:E$902, {TOP20}))*82",
        )
    last = table_end or eg.FIRST_ROW + n - 1
    for col, header in zip(
        ["AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL"], "abcdefgh", strict=True
    ):
        fa_f[f"{col}2"] = header
    fa_f.add_table(Table(displayName=eg.RANKING_TABLE, ref=f"AE2:AL{last}"))
    al = values[eg.ALIASES]
    al.append(["Name1", "Name2"])
    al.append(["Matt Boldy ", "Matthew Boldy"])
    al.append([None, None])
    al.append(["Tim Stutzle", "Tim Stützle"])
    return values, formulas


ROWS = [
    Row("11", "Nathan", "MacKinnon", 70, 0.8, aav=12_600_000),
    Row("22", "Cole", "Caufield", 70, 0.5, aav="#N/A"),
    Row("33", "Tied", "One", 60, 0.3),
    Row("44", "Tied", "Two", 60, 0.3),
    Row("55", "Low", "Gp", 1, 0),
]


def test_extracts_players_divisors_and_aliases() -> None:
    golden = eg.extract(*build(ROWS))
    first = golden.players[0]
    assert first == {
        "player_id": "11",
        "name": "Nathan MacKinnon",
        "team": "COL",
        "position": "C",
        "gp": 70,
        "stats": {"G": 1, "A": 2, "PIM": 3, "PPP": 9, "HIT": 6, "BLK": 7, "SOG": 8},
        "ttltst": 0.8,
        "percentile": pytest.approx(80.0),
        "aav": 12_600_000,
    }
    assert [p["player_id"] for p in golden.players] == ["11", "22", "33", "44", "55"]
    assert golden.divisors == DIVISORS
    assert golden.aliases == [
        {"stats_name": "Matt Boldy", "salary_name": "Matthew Boldy"},
        {"stats_name": "Tim Stutzle", "salary_name": "Tim Stützle"},
    ]
    assert golden.unrated_source_players == 0


def test_unrated_players_have_no_percentile() -> None:
    assert eg.extract(*build(ROWS)).players[4]["percentile"] is None


def test_missing_cap_hit_is_null() -> None:
    assert eg.extract(*build(ROWS)).players[1]["aav"] is None


def test_tied_players_share_the_first_ranks_percentile() -> None:
    players = eg.extract(*build(ROWS)).players
    # Ranks 3 and 4 hold the tie: both get rank 3's percentile.
    assert players[2]["percentile"] == players[3]["percentile"] == pytest.approx(40.0)


def test_rows_follow_their_own_source_row_and_skipped_rows_are_counted() -> None:
    # The last analysis row pulls source row 9, skipping rows 6-8.
    values, formulas = build(ROWS, source_rows=[2, 3, 4, 5, 9])
    golden = eg.extract(values, formulas)
    assert golden.players[-1]["player_id"] == "55"
    assert golden.unrated_source_players == 3


def test_changed_formula_is_a_layout_error() -> None:
    values, formulas = build(ROWS)
    formulas[eg.ANALYSIS]["Q4"] = "=_xlfn.NUMBERVALUE('Stats Data Source'!AD3)"  # hits, not blocks
    with pytest.raises(eg.LayoutError, match="Q4"):
        eg.extract(values, formulas)


def test_changed_divisor_formula_is_a_layout_error() -> None:
    values, formulas = build(ROWS)
    formulas[eg.ANALYSIS]["X6"] = ArrayFormula(
        "X6", f"=AVERAGE(LARGE(F3:F902,{TOP20}))/AVERAGE(LARGE(E$3:E$902, {TOP20}))*82"
    )
    with pytest.raises(eg.LayoutError, match=r"X6 \(A\)"):
        eg.extract(values, formulas)


@pytest.mark.parametrize(
    "formula",
    [
        # end column differs from the start column
        f"=AVERAGE(LARGE(G3:K902,{TOP20}))/AVERAGE(LARGE(E$3:E$902, {TOP20}))*82",
        # not the top 20
        "=AVERAGE(LARGE(G3:G902,{1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1,1}))"
        f"/AVERAGE(LARGE(E$3:E$902, {TOP20}))*82",
        # stat range stops before the last player row (7)
        f"=AVERAGE(LARGE(G3:G6,{TOP20}))/AVERAGE(LARGE(E$3:E$902, {TOP20}))*82",
        # GP range stops before the last player row
        f"=AVERAGE(LARGE(G3:G902,{TOP20}))/AVERAGE(LARGE(E$3:E$5, {TOP20}))*82",
    ],
)
def test_divisor_formula_must_cover_one_column_and_the_top_20(formula: str) -> None:
    values, formulas = build(ROWS)
    formulas[eg.ANALYSIS]["X6"] = ArrayFormula("X6", formula)
    with pytest.raises(eg.LayoutError, match=r"X6 \(A\)"):
        eg.extract(values, formulas)


def test_changed_ranking_formula_is_a_layout_error() -> None:
    values, formulas = build(ROWS)
    formulas[eg.ANALYSIS]["AG5"] = "=(1-AE5/845)*100"
    with pytest.raises(eg.LayoutError, match="AG5"):
        eg.extract(values, formulas)


def test_unreadable_source_row_is_a_layout_error() -> None:
    values, formulas = build(ROWS)
    formulas[eg.ANALYSIS]["A4"] = "Cole Caufield"
    with pytest.raises(eg.LayoutError, match="A4"):
        eg.extract(values, formulas)


def test_two_rows_pulling_one_source_row_is_a_layout_error() -> None:
    with pytest.raises(eg.LayoutError, match="same source row"):
        eg.extract(*build(ROWS, source_rows=[2, 3, 4, 5, 5]))


def test_duplicate_ids_are_a_layout_error() -> None:
    values, formulas = build(ROWS)
    values[eg.STATS]["A3"] = "11"
    with pytest.raises(eg.LayoutError, match="duplicate"):
        eg.extract(values, formulas)


@pytest.mark.parametrize(("cell", "label"), [("A1", "player"), ("AD1", "blocks"), ("AE1", "hits")])
def test_changed_source_header_is_a_layout_error(cell: str, label: str) -> None:
    # Swapped source columns would pass every parity test (each divisor pairs
    # with the same column), so the headers are checked.
    values, formulas = build(ROWS)
    values[eg.STATS][cell] = label
    with pytest.raises(eg.LayoutError, match=cell):
        eg.extract(values, formulas)


def test_ranking_table_must_end_at_the_last_player() -> None:
    # Percentile divides by the table's row count; a longer table skews it.
    with pytest.raises(eg.LayoutError, match="Table2 must span"):
        eg.extract(*build(ROWS, table_end=20))


def test_ranking_must_count_one_per_row() -> None:
    values, formulas = build(ROWS)
    values[eg.ANALYSIS]["AE5"] = 7
    with pytest.raises(eg.LayoutError, match=r"rank 1\.\.5"):
        eg.extract(values, formulas)


@pytest.mark.parametrize(
    ("cell", "bad"), [("E3", "70"), ("F3", 1.5), ("AD3", "#DIV/0!"), ("E3", True)]
)
def test_non_numeric_values_are_layout_errors(cell: str, bad: str | float) -> None:
    values, formulas = build(ROWS)
    values[eg.ANALYSIS][cell] = bad
    with pytest.raises(eg.LayoutError, match=cell):
        eg.extract(values, formulas)


def test_no_player_rows_is_a_layout_error() -> None:
    values, formulas = build(ROWS)
    values[eg.ANALYSIS]["A3"] = None
    with pytest.raises(eg.LayoutError, match="no player rows"):
        eg.extract(values, formulas)


def test_alias_sheet_header_and_rows_are_checked() -> None:
    values, formulas = build(ROWS)
    values[eg.ALIASES]["A1"] = "From"
    with pytest.raises(eg.LayoutError, match="header"):
        eg.extract(values, formulas)
    values, formulas = build(ROWS)
    values[eg.ALIASES]["B2"] = None
    with pytest.raises(eg.LayoutError, match="incomplete"):
        eg.extract(values, formulas)


def test_write_round_trip(tmp_path: Path) -> None:
    golden = eg.extract(*build(ROWS))
    paths = eg.write(golden, tmp_path / "out")
    assert [p.name for p in paths] == [
        "golden_players.json",
        "golden_divisors.json",
        "alias_seed.json",
    ]
    assert json.loads(paths[1].read_text(encoding="utf-8")) == DIVISORS
    assert "Stützle" in paths[2].read_text(encoding="utf-8")  # UTF-8, not \u escapes
    assert len(json.loads(paths[0].read_text(encoding="utf-8"))) == len(ROWS)


def test_main_writes_fixtures(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    golden = eg.extract(*build(ROWS, source_rows=[2, 3, 4, 5, 7]))
    monkeypatch.setattr(eg, "load", lambda path: golden)
    assert eg.main(["book.xlsx", str(tmp_path)]) == 0
    out = capsys.readouterr().out
    assert "5 players, 7 divisors, 2 aliases; 1 Stats Data Source players" in out
    assert (tmp_path / "golden_players.json").is_file()


def test_main_reports_missing_workbook(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    assert eg.main([str(tmp_path / "missing.xlsx"), str(tmp_path)]) == 1
    assert "extract_golden:" in capsys.readouterr().err
    assert not list(tmp_path.iterdir())


def test_main_reports_layout_errors(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    def broken(path: Path) -> eg.Golden:
        raise eg.LayoutError("Fantasy Analysis: unexpected formulas")

    monkeypatch.setattr(eg, "load", broken)
    assert eg.main(["book.xlsx", str(tmp_path)]) == 1
    assert "unexpected formulas" in capsys.readouterr().err


def test_load_reads_a_saved_workbook(tmp_path: Path) -> None:
    # A saved openpyxl workbook has formulas but no cached values, so the
    # value read fails cleanly on the first number rather than guessing.
    _, formulas = build(ROWS)
    path = tmp_path / "book.xlsx"
    formulas.save(path)
    with pytest.raises(eg.LayoutError):
        eg.load(path)
